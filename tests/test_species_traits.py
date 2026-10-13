"""Checks for attribution, taxonomy, uncertainty, IUCN scope and offline replication."""
import importlib.util
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest

import pandas as pd
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
SPEC=importlib.util.spec_from_file_location('traits',ROOT/'Scripts/build_species_traits.py')
M=importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(M)


class SpeciesTraitsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory()
        cls.out=Path(cls.temp.name)
        cls.coverage=M.build(cls.out)
        cls.base=M.csv(ROOT/'Output/LifeHistory/life_history_80_species.csv')
        cls.evidence=M.csv(cls.out/'trait_evidence.csv').set_index('evidence_id')
        cls.iucn=M.csv(cls.out/'iucn_status.csv').set_index('species')
        cls.trophic=M.csv(cls.out/'trophic_traits.csv').set_index('species')

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_original_values_and_roster_preserved(self):
        summary=M.csv(self.out/'species_summary.csv')
        pd.testing.assert_frame_equal(summary[self.base.columns],self.base)
        self.assertEqual(len(summary),80)
        self.assertTrue(summary.species.is_unique)

    def test_ambiguous_synonym_never_inherits_traits(self):
        self.assertFalse((self.evidence.species=='Myctophum opalinum').any())
        self.assertEqual(self.iucn.loc['Myctophum opalinum','verification'],'taxon unresolved')
        self.assertEqual(self.trophic.loc['Myctophum opalinum','trophic_level_preferred'],'')

    def test_iucn_missing_is_not_ne_or_lc_and_latest_not_claimed(self):
        self.assertEqual(len(self.iucn),80)
        self.assertTrue((self.iucn.current_latest_verified=='false').all())
        ne=self.iucn[self.iucn.iucn_raw_code=='N.E.']
        self.assertEqual(len(ne),7)
        self.assertTrue((ne.iucn_category_reported=='').all())
        self.assertEqual((self.iucn.iucn_category_reported!='').sum(),72)
        self.assertEqual(self.iucn.loc['Glyptocephalus cynoglossus','assessment_id'],'162704857')
        self.assertEqual(self.iucn.loc['Glyptocephalus cynoglossus','assessment_scope'],'Global')
        self.assertEqual(self.iucn.loc['Glyptocephalus cynoglossus','assessment_review_evidence_ids'],'CUR012')

    def test_iucn_population_stocks_not_substituted(self):
        herring=self.iucn.loc['Clupea harengus']
        self.assertEqual(herring.iucn_category_reported,'LC')
        self.assertEqual(herring.assessment_id,'4717767')
        evidence=self.evidence.loc[herring.evidence_id]
        self.assertEqual(json.loads(evidence.context_json)['Level'],'species in general')

    def test_trophic_methods_and_uncertainty_stay_paired(self):
        for sp,r in self.trophic.iterrows():
            if not r.trophic_level_preferred:
                continue
            suffix='diet' if r.trophic_level_diet else 'food_items'
            self.assertEqual(r.trophic_level_method,suffix)
            self.assertEqual(r.trophic_level_preferred,r['trophic_level_'+suffix])
            self.assertEqual(r.trophic_level_se,r['trophic_level_'+suffix+'_se'])
            e=self.evidence.loc[r.trophic_level_evidence_id]
            self.assertEqual(e.species,sp)
            self.assertEqual(e.evidence_type,'derived_trophic_estimate')
            self.assertEqual(float(e.value),float(r.trophic_level_preferred))

    def test_sources_and_summary_links_resolve(self):
        sources=M.csv(self.out/'sources.csv').set_index('source_id')
        for c in ['citation','url','sha256','retrieved_at']:
            self.assertTrue((sources[c]!='').all())
        self.assertTrue(set(self.evidence.source_id)<=set(sources.index))
        links=M.csv(self.out/'summary_evidence.csv')
        for r in links.itertuples():
            e=self.evidence.loc[r.evidence_id]
            self.assertEqual(e.species,r.species)
            self.assertTrue(e.locator)
            self.assertNotEqual(e.decision,'hold')

    def test_reef_limits_and_disagreements_retained(self):
        e=self.evidence.loc['REEF|Dibranchus atlanticus|bodysize_max']
        self.assertEqual(e.value,'1')
        self.assertEqual(e.decision,'hold')
        self.assertIn('large_disagreement',e.qc_flags)
        e=self.evidence.loc['REEF|Caranx latus|rafter']
        self.assertIn('genus_level',e.qc_flags)
        pinfish=self.trophic.loc['Lagodon rhomboides']
        self.assertIn('ontogenetic',pinfish.source_comparison)
        self.assertIn('adult omnivory',pinfish.contextual_diet_guild)
        issues=M.csv(self.out/'source_issues.csv')
        self.assertEqual(len(issues),2)
        self.assertTrue(issues.source_species.str.contains('Cerdale').all())

    def test_contextual_estimates_preserve_population_and_bounds(self):
        generation=self.evidence.loc['CUR001']
        self.assertEqual(generation.value_min,'16')
        self.assertEqual(generation.value_max,'17')
        self.assertEqual(generation.evidence_type,'assessment_derived_estimate')
        self.assertEqual(self.evidence.loc['CUR002','value_max'],'1')
        self.assertEqual(self.evidence.loc['CUR002','unit'],'years')
        contexts=[json.loads(self.evidence.loc[k,'context_json']) for k in ['CUR003','CUR004','CUR005']]
        self.assertNotEqual(contexts[0]['population'],contexts[1]['population'])
        self.assertNotEqual(contexts[1]['sex'],contexts[2]['sex'])

    def test_gap_inventory_covers_every_species_and_domain(self):
        gaps=M.csv(self.out/'trait_gaps.csv')
        self.assertEqual(len(gaps),80*len(M.DOMAINS))
        self.assertFalse(gaps.duplicated(['species','trait_domain']).any())
        self.assertTrue((gaps.groupby('trait_domain').species.nunique()==80).all())
        self.assertIn('larval_duration',set(gaps.trait_domain))
        self.assertIn('fecundity',set(gaps.trait_domain))
        self.assertIn('model_only',set(gaps.coverage_status))
        self.assertGreater((gaps.loc[gaps.trait_domain=='feeding_mechanism','coverage_status']=='available_contextual').sum(),60)
        fields=M.csv(self.out/'trait_field_gaps.csv')
        self.assertFalse(fields.duplicated(['species','trait']).any())
        self.assertTrue((fields.groupby('trait').species.nunique()==80).all())

    def test_sqlite_integrity_and_numeric_types(self):
        with sqlite3.connect(self.out/'species_traits.sqlite') as db:
            self.assertEqual(db.execute('PRAGMA integrity_check').fetchone()[0],'ok')
            self.assertFalse(db.execute('PRAGMA foreign_key_check').fetchall())
            self.assertEqual(db.execute('SELECT count(*) FROM species').fetchone()[0],80)
            self.assertEqual(db.execute("SELECT value_min, value_max FROM trait_evidence WHERE evidence_id='CUR001'").fetchone(),(16.0,17.0))
            self.assertEqual(db.execute("SELECT typeof(trophic_level_preferred) FROM trophic_traits WHERE trophic_level_preferred IS NOT NULL LIMIT 1").fetchone()[0],'real')
            self.assertGreater(db.execute('SELECT count(*) FROM trait_provenance WHERE citation IS NOT NULL').fetchone()[0],2600)

    def test_excel_escape_is_explicit_and_canonical_value_preserved(self):
        f=pd.DataFrame({'value':['a\x02b','x'*33000]})
        changes=[]
        safe=M.excel_frame(f,'example',changes)
        self.assertEqual(f.loc[0,'value'],'a\x02b')
        self.assertEqual(safe.loc[0,'value'],'a\\u0002b')
        self.assertLessEqual(len(safe.loc[1,'value']),32767)
        self.assertEqual(len(changes),2)
        w=openpyxl.load_workbook(self.out/'species_traits_80.xlsx',read_only=True)
        self.assertEqual(w['species_summary'].max_row,81)
        self.assertIn('excel_export_notes',w.sheetnames)
        w.close()

    def test_repeated_offline_build_has_identical_canonical_outputs(self):
        paths=list(self.out.glob('*.csv'))+[self.out/'coverage.json']
        before={p.name:M.sha(p) for p in paths}
        M.build(self.out)
        self.assertEqual(before,{p.name:M.sha(p) for p in paths})


if __name__=='__main__':
    unittest.main()
