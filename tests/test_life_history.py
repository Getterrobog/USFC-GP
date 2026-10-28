"""Scientific join/provenance checks for the populated trait extraction."""
import importlib.util
import json
from pathlib import Path
import unittest

import openpyxl
import pandas as pd
import pyarrow.parquet as pq

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('life_history', ROOT/'Scripts/build_life_history.py')
lh = importlib.util.module_from_spec(spec)
spec.loader.exec_module(lh)


class LifeHistoryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wide = pd.read_csv(lh.OUT/'life_history_80_species.csv', keep_default_na=False)
        cls.evidence = pd.read_csv(lh.OUT/'trait_evidence.csv', keep_default_na=False)
        cls.roster = pd.read_csv(lh.ROSTER)
        cls.cross = pd.read_csv(lh.CROSSWALK, keep_default_na=False)
        cls.species = pq.read_table(lh.RAW/'species.parquet').to_pandas(timestamp_as_object=True)
        cls.synonyms = pq.read_table(lh.RAW/'synonyms.parquet').to_pandas(timestamp_as_object=True)

    def test_frozen_roster_no_subspecies_expansion(self):
        self.assertEqual(self.wide.species.tolist(), self.roster.species.tolist())
        self.assertEqual(len(self.wide), 80)
        self.assertTrue(self.wide.species.is_unique)
        self.assertTrue(self.wide.species.str.split().map(len).eq(2).all())

    def test_uncertain_synonym_never_copies_affine_traits(self):
        row = self.wide.set_index('species').loc['Myctophum opalinum']
        self.assertEqual(row.fishbase_spec_code, '')
        self.assertEqual(row.generation_time_status, 'taxon_match_withheld')
        self.assertFalse((self.evidence.species == 'Myctophum opalinum').any())
        self.assertTrue((self.evidence.species == 'Myctophum affine').any())

    def test_ambiguous_override_fails_closed(self):
        cross = self.cross.copy()
        cross.loc[cross.species == 'Myctophum opalinum', 'attach_traits'] = 1
        with self.assertRaisesRegex(ValueError, 'Ambiguous synonym'):
            lh.match_taxa(self.roster, self.species, self.synonyms, cross)

    def test_unreviewed_name_fails_closed(self):
        roster = self.roster.copy()
        roster.loc[0, 'species'] = 'Unreviewed taxon'
        with self.assertRaisesRegex(ValueError, 'Unreviewed unmatched'):
            lh.match_taxa(roster, self.species, self.synonyms, self.cross)

    def test_no_duplicate_evidence_and_references_are_auditable(self):
        self.assertTrue(self.evidence.evidence_id.is_unique)
        refs = pd.read_csv(lh.OUT/'references.csv', keep_default_na=False)
        ids = set(refs.RefNo.astype(str))
        review = pd.read_csv(lh.OUT/'review_queue.csv', keep_default_na=False)
        for r in self.evidence.itertuples():
            if r.selected_reference_id and str(r.selected_reference_id) not in ids:
                self.assertTrue(((review.evidence_id == r.evidence_id) &
                                 (review.reason == 'reference_id_not_in_snapshot')).any())
        self.assertTrue(self.evidence.source_sha256.str.len().eq(64).all())

    def test_model_not_reported_and_no_imputed_zero(self):
        model = self.evidence[self.evidence.trait == 'generation_time_model']
        self.assertTrue(model.evidence_type.eq('model_or_derived_keyfacts').all())
        self.assertTrue(pd.to_numeric(model.value).gt(0).all())
        missing = self.wide.generation_time_status.ne('model_or_derived_keyfacts')
        self.assertTrue(self.wide.loc[missing, 'generation_time_model_years'].eq('').all())

    def test_reversed_larval_bounds_not_repaired_or_summarized(self):
        larval = self.evidence[self.evidence.trait.str.startswith('LarvalDuration')]
        self.assertEqual(larval.species.nunique(), 3)
        self.assertTrue(larval.qc_flags.str.contains('source_min_exceeds_max').all())
        self.assertTrue(larval.unit.eq('source_duration_unit_unverified').all())
        lutjanus = larval[larval.species == 'Lutjanus synagris'].set_index('source_field')
        self.assertEqual(float(lutjanus.loc['LarvalDurationMin', 'value']), 40)
        self.assertEqual(float(lutjanus.loc['LarvalDurationMax', 'value']), 15)

    def test_maturity_record_envelope_preserves_sex_context(self):
        e = self.evidence[(self.evidence.species == 'Alosa aestivalis') &
                          (self.evidence.trait == 'maturity_age_reported')]
        self.assertEqual({json.loads(v)['Sex'] for v in e.source_context_json}, {'male','female'})
        r = self.wide.set_index('species').loc['Alosa aestivalis']
        self.assertEqual(float(r.maturity_age_reported_years_min), 3.5)
        self.assertEqual(float(r.maturity_age_reported_years_max), 4)

    def test_fecundity_type_is_retained_and_unspecified_flagged(self):
        f = self.evidence[self.evidence.source_table == 'fecundity']
        for r in f.itertuples():
            if 'FecundityType' not in json.loads(r.source_context_json):
                self.assertIn('fecundity_basis_unspecified', r.qc_flags)
        self.assertFalse(any(c in self.wide for c in ['fecundity_min','fecundity_max']))

    def test_workbook_roundtrip_and_field_dictionary(self):
        book = openpyxl.load_workbook(lh.OUT/'life_history_80_species.xlsx', read_only=True)
        self.assertEqual(book['80 species'].max_row, 81)
        self.assertEqual(book['Evidence'].max_row, len(self.evidence)+1)
        self.assertEqual(list(book['80 species'].values)[1][0], self.roster.species.iloc[0])
        book.close()
        d = pd.read_csv(lh.OUT/'table_dictionary.csv')
        self.assertEqual(d.field.tolist(), self.wide.columns.tolist())

    def test_all_archived_source_hashes(self):
        manifest = json.loads((lh.RAW/'download_manifest.json').read_text())
        for name, record in manifest.items():
            self.assertEqual(lh.sha(lh.RAW/f'{name}.parquet'), record['sha256'])


if __name__ == '__main__':
    unittest.main()
