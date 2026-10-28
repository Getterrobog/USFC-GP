"""Checks for image identity joins, source licenses and missing-data behavior."""
import importlib.util
import json
from pathlib import Path
import unittest

import openpyxl
import pandas as pd

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('images',ROOT/'Scripts/build_fish_image_catalog.py')
images=importlib.util.module_from_spec(spec)
spec.loader.exec_module(images)


class ImageCatalogTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rows=json.loads((images.OUT/'fish_images_80_species.json').read_text())
        cls.by_name={r['species']:r for r in cls.rows}
        cls.candidates=pd.read_csv(images.OUT/'image_candidates.csv',keep_default_na=False)

    def test_all_project_species_retained_once(self):
        roster=pd.read_csv(images.ROSTER)
        self.assertEqual([r['species'] for r in self.rows],roster.species.tolist())
        self.assertEqual(len(self.by_name),80)

    def test_selected_urls_have_image_responses(self):
        selected=[r for r in self.rows if r['image_url']]
        self.assertEqual(len(selected),78)
        for r in selected:
            self.assertEqual(r['image_status'],'verified_image_response')
            self.assertEqual(r['image_http_status'],200)
            self.assertTrue(r['image_content_type'].startswith('image/'))
            self.assertEqual(len(r['image_sha256']),64)
            self.assertTrue(r['page_image_match'])

    def test_selections_are_eligible_whole_fish_images(self):
        selected=self.candidates[self.candidates.selected]
        self.assertEqual(len(selected),78)
        self.assertTrue(selected.eligible.all())
        self.assertTrue(selected.life_stage.isin(['male','female','unsexed']).all())
        self.assertFalse(selected.image_type.isin(['stamp','fish details']).any())

    def test_no_uncertain_synonym_or_family_placeholder_substitution(self):
        for name in ['Myctophum opalinum','Prionotus martis']:
            self.assertEqual(self.by_name[name]['image_url'],'')
            self.assertFalse(self.by_name[name]['eligible_for_noncommercial_embedding'])
        self.assertEqual(self.by_name['Myctophum opalinum']['fishbase_spec_code'],'')

    def test_published_eligibility_requires_verified_picture_license(self):
        for r in self.rows:
            if r['eligible_for_noncommercial_embedding']:
                self.assertTrue(r['credit'])
                self.assertTrue(r['license_url'])
                self.assertTrue(r['page_image_match'])
                self.assertIn(r['reuse_status'],['attribution_required','attribution_noncommercial'])
                self.assertEqual(r['image_status'],'verified_image_response')
            if r['reuse_status']=='permission_required':
                self.assertFalse(r['eligible_for_noncommercial_embedding'])

    def test_license_parser_rejects_wrong_picture_and_generic_footer(self):
        license_link="<a rel='license' href='https://creativecommons.org/licenses/by-nc/4.0/'>CC BY-NC</a>"
        wrong="<img src='../images/species/wrong.jpg'>"+license_link+'Image can be freely copied'
        self.assertEqual(images.page_rights(wrong,'fish.jpg')['reuse_status'],'page_image_mismatch')
        footer="<img src='../images/species/fish.jpg'>"+license_link
        self.assertEqual(images.page_rights(footer,'fish.jpg')['reuse_status'],'license_unverified')
        valid=footer+'Image can be freely copied'
        self.assertEqual(images.page_rights(valid,'fish.jpg')['reuse_status'],'attribution_noncommercial')
        self.assertEqual(images.page_rights(valid+' ALL RIGHTS RESERVED','fish.jpg')['reuse_status'],'permission_required')

    def test_alternate_leads_stay_separate_and_flagged(self):
        alternates=[r for r in self.rows if r['supplemental_image_url']]
        self.assertEqual(len(alternates),2)
        self.assertTrue(all(r['supplemental_image_status']=='verified_image_response' for r in alternates))
        self.assertTrue(all(not r['eligible_for_noncommercial_embedding'] for r in alternates))
        self.assertEqual(self.by_name['Myctophum opalinum']['supplemental_review_status'],'taxonomic_review_required')
        self.assertEqual(self.by_name['Prionotus martis']['supplemental_reuse_status'],'permission_required')

    def test_source_page_integrity_and_picture_id_url(self):
        checks=json.loads(images.CHECKS.read_text())
        for check in checks.values():
            self.assertEqual(images.sha((ROOT/check['page_cache']).read_bytes()),check['page_sha256'])
        _,url=images.urls('1574.0','Alaes_u0.gif')
        self.assertIn('ID=1574&',url)
        self.assertIn('pic=Alaes_u0.gif',url)

    def test_workbook_is_readable_and_contains_all_rows(self):
        book=openpyxl.load_workbook(images.OUT/'fish_images_80_species.xlsx',read_only=True)
        self.assertEqual(book['80 species'].max_row,81)
        self.assertEqual(book['All candidates'].max_row,len(self.candidates)+1)
        book.close()


if __name__=='__main__':
    unittest.main()
