#!/usr/bin/env python3
"""Build a traceable, offline trait extraction for the frozen 80-species roster.

No taxonomy cleaning, fuzzy joins, trait imputation, or ranking is performed.
See docs/life_history_workflow.md for the scientific scope and review limitations.
"""
from __future__ import annotations
import hashlib
import json
import platform
import re
from pathlib import Path
import subprocess

import numpy as np
import pandas as pd
import pyarrow
import pyarrow.parquet as pq
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'Data/LifeHistory/raw/v26.06'
OUT = ROOT / 'Output/LifeHistory'
ROSTER = ROOT / 'Output/SpeciesSelection/species_investigation_queue.csv'
CROSSWALK = ROOT / 'config/life_history_taxon_crosswalk.csv'
VERSION = '26.06'
KEYFACTS = 'https://www.fishbase.se/manual/key%20facts.htm'

# table, field, trait, unit, direct reference field, evidence type, length type
# "reported" means transcribed from FishBase, not independently verified or
# necessarily measured directly: population growth fits are estimates too.
FIELDS = []
def add(table, field, trait, unit, ref='', kind='reported_database_value', length=''):
    FIELDS.append(dict(table=table, field=field, trait=trait, unit=unit,
                       reference_field=ref, evidence_type=kind, length_type_field=length))

for field, trait, unit, ref, lt in [
    ('Length','maximum_length_male_unsexed','cm','MaxLengthRef','LTypeMaxM'),
    ('LengthFemale','maximum_length_female','cm','MaxLengthRef','LTypeMaxF'),
    ('LongevityWild','maximum_age_wild','years','LongevityWildRef',''),
    ('LongevityCaptive','maximum_age_captive','years','LongevityCapRef',''),
    ('DepthRangeShallow','depth_shallow','m','DepthRangeRef',''),
    ('DepthRangeDeep','depth_deep','m','DepthRangeRef',''),
    ('DemersPelag','water_column_habitat','category','',''),
    ('AnaCat','migration','category','MigratRef',''),
    ('Importance','fisheries_importance','category','ImportanceRef','')]:
    add('species',field,trait,unit,ref,length=lt)
for f in ['Fresh','Brack','Saltwater']:
    add('species',f,'water_type_'+f.lower(),'source_flag',kind='database_classification')
for field, trait, unit in [
    ('Generation_time','generation_time_model','years'),
    ('Gen_time_1st','generation_time_model_lower','years'),
    ('Gen_time_2nd','generation_time_model_upper','years'),
    ('Life_span','lifespan_model','years'),
    ('tm','maturity_age_model','years'),
    ('K','growth_k_keyfacts','1/year')]:
    add('matrix',field,trait,unit,kind='model_or_derived_keyfacts')
for field,trait,ref,unit,lt in [
    ('tm','maturity_age_reported','AgeMatRef','years',''),
    ('AgeMatMin','maturity_age_source_min','AgeMatRef','years',''),
    ('AgeMatMin2','maturity_age_source_second','AgeMatRef','years',''),
    ('Lm','maturity_length_reported','LengthMatRef','cm','Type1'),
    ('LengthMatMin','maturity_length_source_min','LengthMatRef','cm','Type1'),
    ('LengthMatMin2','maturity_length_source_second','LengthMatRef','cm','Type1')]:
    add('maturity',field,trait,unit,ref,length=lt)
for field,trait,ref,unit in [
    ('K','growth_k_reported','DataSourceRef','1/year'),
    ('tmax','maximum_age_growth_record','tmaxRef','years'),
    ('tm','maturity_age_growth_record','DataSourceRef','years')]:
    add('popgrowth',field,trait,unit,ref)
for field,ref in [('ReproMode','ReproducRefNo'),('Fertilization','ReproducRefNo'),
                  ('RepGuild1','ReproducRefNo'),('RepGuild2','ReproducRefNo'),
                  ('ParentalCare','ParentalCareRef'),('Spawning','SpawnModRef')]:
    add('reproduc',field,field,'category',ref)
for f in ['FecundityMin','FecundityMax','FecundityMean']:
    add('fecundity',f,f,'eggs_or_offspring_per_female; basis_in_FecundityType','FecundityRef')
for f in ['LarvalDurationMin','LarvalDurationMax','LarvalDurationMod']:
    add('larvae',f,f,'source_duration_unit_unverified',f+'Ref')
for f,r in [('PlaceofDevelopment','LarvaeRefNo'),('LarvalArea','LarvaeRefNo')]:
    add('larvae',f,f,'category',r)
for f,r in [('DietTroph','DietRef'),('FoodTroph','FoodRef')]:
    add('ecology',f,f,'trophic_level',r)
add('ecology','FeedingType','feeding_type','category','FeedingTypeRef')
HABITATS = ['Neritic','Oceanic','Epipelagic','Mesopelagic','Bathypelagic',
            'Estuaries','Mangroves','Saltmarshes','Intertidal','Stream','Lakes',
            'Sand','Mud','Rocky','SeaGrassBeds','CoralReefs']
for f in HABITATS:
    ref = 'SubstrateRef' if f in ['Sand','Mud','Rocky'] else (
        'SpecialHabitatRef' if f in ['SeaGrassBeds','CoralReefs'] else 'HabitatsRef')
    add('ecology',f,'habitat_flag_'+f,'source_flag',ref,kind='database_classification')

MAIN_REFS = dict(species='SpeciesRefNo', maturity='MaturityRefNo',
                 popgrowth='PopGrowthRef', reproduc='ReproducRefNo',
                 fecundity='MainRefNo', larvae='LarvaeRefNo', ecology='EcologyRefNo')
IDS = dict(species='SpecCode',matrix='ID',maturity='autoctr',popgrowth='AutoCtr',
           reproduc='autoctr',fecundity='autoctr',larvae='autoctr',ecology='autoctr')
CONTEXT = ['StockCode','stockcode','Sex','Locality','C_Code','E_CODE','YearStart','YearEnd',
           'GrowthEnviron','MethodGrowth','Data','FecundityType','LengthFecunMin',
           'LengthFecunMax','LengthFecunMean','LengthTypeFecMin','LengthTypeFecMax',
           'LengthTypeFecMean','WeightMin','WeightMax','WeightMean','FecComment',
           'Comment','Comment2','Type1','LCL_tm','UCL_tm','SE_tm','SD_tm',
           'LCL_Lm','UCL_Lm','LCL_K','UCL_K','SE_K','SD_K','Number','KF_type']


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def present(value):
    return value is not None and not pd.isna(value) and str(value).strip() != ''


def clean(value):
    if not present(value):
        return None
    if hasattr(value, 'item'):
        value = value.item()
    return value


def ref_id(value):
    if not present(value) or float(value) <= 0:
        return ''
    return str(int(value))


def match_taxa(roster, species, synonyms, exceptions):
    """Fail closed on new names, ambiguous overrides, and source-key drift."""
    if len(roster) != 80 or not roster.species.is_unique:
        raise ValueError('Expected the frozen 80 unique project species')
    if not exceptions.species.is_unique:
        raise ValueError('Duplicate crosswalk exception')
    species = species.assign(source_name=species.Genus + ' ' + species.Species)
    if not species.source_name.is_unique:
        raise ValueError('Duplicate source accepted names')
    exact = species.set_index('source_name')
    ex = exceptions.set_index('species')
    rows = []
    for name in roster.species:
        if name in exact.index:
            r = dict(species=name,fishbase_spec_code=int(exact.loc[name,'SpecCode']),
                     fishbase_name=name,syn_code=None,match_status='exact_accepted_name',
                     attach_traits=1,decision_basis='Exact match to pinned species table',evidence_url='')
        else:
            if name not in ex.index:
                raise ValueError(f'Unreviewed unmatched name: {name}')
            r = dict(species=name, **ex.loc[name].to_dict())
            source = species[species.SpecCode == int(r['fishbase_spec_code'])]
            if len(source) != 1 or source.iloc[0].source_name != r['fishbase_name']:
                raise ValueError(f'Source crosswalk drift: {name}')
            syn = synonyms[synonyms.SynCode == int(r['syn_code'])]
            if len(syn) != 1 or int(syn.iloc[0].SpecCode) != int(r['fishbase_spec_code']):
                raise ValueError(f'Synonym crosswalk drift: {name}')
            s = syn.iloc[0]
            if s.SynGenus + ' ' + s.SynSpecies != name:
                raise ValueError(f'Synonym name mismatch: {name}')
            if int(r['attach_traits']) and s.Status != 'synonym':
                raise ValueError(f'Ambiguous synonym cannot transfer traits: {name}')
            r['synonym_reference_id'] = ref_id(s.SynonymsRef)
            r['synonym_authority'] = s.Author
        r['attached_spec_code'] = int(r['fishbase_spec_code']) if int(r['attach_traits']) else None
        rows.append(r)
    out = pd.DataFrame(rows)
    attached = out.attached_spec_code.dropna()
    if not attached.is_unique:
        raise ValueError('Multiple project species share an attached FishBase taxon; review required')
    return out


def quality_flags(table, row, spec):
    flags = []
    f = spec['field']
    value = row.get(f)
    if spec['unit'] not in ['category','source_flag'] and present(value):
        if float(value) < 0:
            flags.append('negative_value')
        if float(value) == 0 and spec['unit'] in ['years','cm','1/year']:
            flags.append('nonpositive_value')
    if table == 'larvae' and f.startswith('LarvalDuration'):
        flags.append('duration_unit_and_definition_unverified')
        if present(row.LarvalDurationMin) and present(row.LarvalDurationMax):
            if row.LarvalDurationMin > row.LarvalDurationMax:
                flags.append('source_min_exceeds_max')
    if table == 'fecundity' and not present(row.FecundityType):
        flags.append('fecundity_basis_unspecified')
    if table == 'fecundity' and present(row.FecundityMin) and present(row.FecundityMax):
        if row.FecundityMin > row.FecundityMax:
            flags.append('source_min_exceeds_max')
    if table == 'matrix' and f.startswith(('Generation','Gen_time')):
        lo, hi = row.Gen_time_1st, row.Gen_time_2nd
        if present(lo) and present(hi) and lo > hi:
            flags.append('source_min_exceeds_max')
        if present(lo) and present(hi) and present(row.Generation_time):
            if not lo <= row.Generation_time <= hi:
                flags.append('estimate_outside_source_bounds')
    if spec['length_type_field'] and not present(row.get(spec['length_type_field'])):
        flags.append('length_type_missing')
    return ';'.join(flags)


def numeric_summary(evidence, trait):
    x = evidence[(evidence.trait == trait) & (evidence.qc_flags == '')]
    values = pd.to_numeric(x.value, errors='raise')
    return (values.min() if len(x) else None, values.max() if len(x) else None, len(x))


def run():
    OUT.mkdir(parents=True, exist_ok=True)
    source_dir = OUT/'source_records'
    source_dir.mkdir(exist_ok=True)
    download = json.loads((RAW/'download_manifest.json').read_text())
    tables = {}
    for table, meta in sorted(download.items()):
        path = RAW/f'{table}.parquet'
        if meta['version'] != VERSION or sha(path) != meta['sha256']:
            raise ValueError(f'Raw snapshot failed checksum/version: {table}')
        tables[table] = pq.read_table(path).to_pandas(timestamp_as_object=True)
    tables['maturity'] = tables['maturity'].rename(columns={'Speccode':'SpecCode'})
    roster = pd.read_csv(ROSTER, keep_default_na=False)
    exceptions = pd.read_csv(CROSSWALK, keep_default_na=False)
    cross = match_taxa(roster, tables['species'], tables['synonyms'], exceptions)
    cross.to_csv(OUT/'taxon_crosswalk.csv',index=False)
    codes = cross.attached_spec_code.dropna().astype(int).tolist()
    # All fields of each matched source row are archived; no population/sex row
    # is chosen as representative. The upstream full parquet is untouched.
    for table, frame in tables.items():
        if table == 'refrens':
            continue
        if table == 'synonyms':
            frame = frame[frame.SynCode.isin(cross.syn_code.dropna().astype(int))]
        else:
            frame = frame[frame.SpecCode.isin(codes)]
        frame.to_csv(source_dir/f'{table}.csv',index=False)
        tables[table] = frame
    evidence = []
    for mapping in cross.to_dict('records'):
        if not mapping['attach_traits']:
            continue
        name, code = mapping['species'], mapping['attached_spec_code']
        for spec in FIELDS:
            table = spec['table']
            for _, row in tables[table][tables[table].SpecCode == code].iterrows():
                value = row[spec['field']]
                if not present(value):
                    continue
                # Source flags are sparse positive evidence; unchecked fields
                # must not be treated as demonstrated ecological absence.
                if spec['unit'] == 'source_flag' and value not in [1,-1]:
                    continue
                direct = ref_id(row.get(spec['reference_field'])) if spec['reference_field'] else ''
                main = ref_id(row.get(MAIN_REFS.get(table,'')))
                rid = str(int(row[IDS[table]]))
                context = {k:clean(row[k]) for k in CONTEXT if k in row and present(row[k])}
                reference = direct or main
                evidence.append(dict(
                    evidence_id=f'{name}|{table}|{rid}|{spec["field"]}',
                    species=name, fishbase_name=mapping['fishbase_name'],
                    fishbase_spec_code=int(code), taxon_match_status=mapping['match_status'],
                    trait=spec['trait'],value=clean(value),unit=spec['unit'],
                    length_type=clean(row.get(spec['length_type_field'])) if spec['length_type_field'] else None,
                    evidence_type=spec['evidence_type'],source_table=table,
                    source_field=spec['field'],source_row_id=rid,
                    direct_reference_id=direct,record_reference_id=main,
                    selected_reference_id=reference,
                    reference_scope='field' if direct else ('record_fallback' if main else 'not_supplied'),
                    reference_url=f'https://www.fishbase.se/references/FBRefSummary.php?ID={reference}' if reference else '',
                    species_url=f'https://www.fishbase.se/summary/{int(code)}',
                    method_url=KEYFACTS if table=='matrix' else '',
                    source_context_json=json.dumps(context,ensure_ascii=False,sort_keys=True),
                    qc_flags=quality_flags(table,row,spec),
                    primary_source_review='not_independently_reviewed',
                    fishbase_version=VERSION,source_sha256=download[table]['sha256'],
                    source_url=download[table]['url'],retrieved_utc=download[table]['retrieved_utc']))
    ev = pd.DataFrame(evidence).sort_values(['species','source_table','source_row_id','source_field'])
    if not ev.evidence_id.is_unique:
        raise ValueError('Duplicate source evidence key')
    ev.to_csv(OUT/'trait_evidence.csv',index=False)
    pd.DataFrame(FIELDS).to_csv(OUT/'trait_dictionary.csv',index=False)

    # Every reference field in selected source records is retained in the
    # bibliography, including context references not used as the primary link.
    references = set()
    for table, frame in tables.items():
        if table == 'refrens':
            continue
        for col in frame:
            if re.search(r'ref(?:no)?[0-9]*$',col,re.I) and pd.api.types.is_numeric_dtype(frame[col]):
                references.update(ref_id(v) for v in frame[col].dropna())
    references.discard('')
    refs = tables['refrens']
    refs = refs[refs.RefNo.isin([int(v) for v in references])]
    ref_columns = ['RefNo','Author','Year','Title','Source','ShortCitation','RefType','DOI','PaperURL']
    refs = refs[ref_columns].sort_values('RefNo').copy()
    refs['fishbase_reference_url'] = refs.RefNo.map(lambda v:f'https://www.fishbase.se/references/FBRefSummary.php?ID={int(v)}')
    refs['primary_source_review'] = 'not_independently_reviewed'
    refs.to_csv(OUT/'references.csv',index=False)

    wide = []
    dictionary = {}
    def describe(field, unit, definition):
        dictionary[field] = dict(field=field,unit=unit,definition=definition)
    # Single-source values preserve sex and length type; never average TL/SL/FL.
    single = {
        'generation_time_model_years':'generation_time_model',
        'generation_time_model_lower':'generation_time_model_lower',
        'generation_time_model_upper':'generation_time_model_upper',
        'lifespan_model_years':'lifespan_model',
        'maturity_age_model_years':'maturity_age_model',
        'maximum_age_wild_years':'maximum_age_wild',
        'maximum_age_captive_years':'maximum_age_captive',
        'maximum_length_male_unsexed_cm':'maximum_length_male_unsexed',
        'maximum_length_female_cm':'maximum_length_female',
        'depth_shallow_m':'depth_shallow','depth_deep_m':'depth_deep',
        'water_column_habitat':'water_column_habitat',
        'migration':'migration','fisheries_importance':'fisheries_importance'}
    aggregates = {'maturity_age_reported_years':'maturity_age_reported',
                  'maturity_age_growth_record_years':'maturity_age_growth_record',
                  'growth_k_reported_per_year':'growth_k_reported',
                  'maximum_age_growth_record_years':'maximum_age_growth_record'}
    categorical = {'sexual_system':'ReproMode','fertilization':'Fertilization',
                   'reproductive_guild_1':'RepGuild1','reproductive_guild_2':'RepGuild2',
                   'parental_care':'ParentalCare','spawning_pattern':'Spawning',
                   'feeding_type':'feeding_type','larval_development_place':'PlaceofDevelopment'}
    for mapping in cross.to_dict('records'):
        name=mapping['species']
        e = ev[ev.species == name]
        selected = roster[roster.species == name].iloc[0]
        row = dict(species=name,investigation_tier=selected.investigation_tier,
                   fishbase_name=mapping['fishbase_name'] if mapping['attach_traits'] else '',
                   fishbase_spec_code=mapping['attached_spec_code'],
                   taxon_match_status=mapping['match_status'])
        for field, trait in single.items():
            x=e[(e.trait==trait) & (e.qc_flags=='')]
            if len(x)>1:
                raise ValueError(f'Multiple records for single-value field {name} {trait}')
            row[field] = x.iloc[0].value if len(x) else None
            spec = next(s for s in FIELDS if s['trait']==trait)
            describe(field,spec['unit'],f'{spec["table"]}.{spec["field"]}; {spec["evidence_type"]}. Source bounds have unverified interval level; not assumed 95% CI.')
            if 'maximum_length' in field:
                k=field.replace('_cm','_type')
                row[k]=x.iloc[0].length_type if len(x) else None
                describe(k,'category','Length convention supplied by FishBase: TL/SL/FL retained without conversion.')
        for prefix, trait in aggregates.items():
            lo,hi,n=numeric_summary(e,trait)
            for suffix,value in [('min',lo),('max',hi),('n_records',n)]:
                k=prefix+'_'+suffix
                row[k]=value
                describe(k,'record_count' if suffix=='n_records' else ('1/year' if 'per_year' in prefix else 'years'),
                    f'{suffix} across unflagged {trait} source rows; discovery envelope across populations and sexes, not a CI or pooled estimate. Counts are database records, not independent studies.')
        for field,trait in categorical.items():
            row[field]='; '.join(sorted(set(str(v) for v in e.loc[e.trait==trait,'value']))) or None
            describe(field,'category','Distinct '+trait+' categories across all matched records; combinations are not a single population profile.')
        row['positive_habitat_tags']='; '.join(sorted(set(e.loc[e.trait.str.startswith('habitat_flag_'),'source_field']))) or None
        row['positive_water_types']='; '.join(sorted(set(e.loc[e.trait.str.startswith('water_type_'),'source_field']))) or None
        lm=e[e.trait=='maturity_length_reported']
        row['maturity_length_n_records']=len(lm)
        row['maturity_length_records']=' | '.join(
            f'{r.value} cm {r.length_type or "type unspecified"} [record {r.source_row_id}; {r.qc_flags or "no automated flag"}]'
            for r in lm.itertuples()) or None
        f=e[e.source_table=='fecundity']
        row['fecundity_n_records']=f.source_row_id.nunique()
        # Preserve each range/type rather than mix batch, annual and unknown bases.
        descriptions=[]
        for rid,g in f.groupby('source_row_id',sort=True):
            context=json.loads(g.iloc[0].source_context_json)
            vals=', '.join(f'{v.source_field}={v.value}' for v in g.itertuples())
            descriptions.append(f'{context.get("FecundityType","basis unspecified")}: {vals} [record {rid}]')
        row['fecundity_records']=' | '.join(descriptions) or None
        duration=e[e.trait.str.startswith('LarvalDuration')]
        row['larval_duration_n_records']=duration.source_row_id.nunique()
        row['larval_duration_status']='needs_source_review' if len(duration) else 'not_found_in_snapshot'
        row['generation_time_status']='model_or_derived_keyfacts' if present(row['generation_time_model_years']) else ('taxon_match_withheld' if not mapping['attach_traits'] else 'not_found_in_snapshot')
        row['primary_source_review']='not_independently_reviewed'
        row['regional_applicability']='not_assessed'
        row['evidence_rows']=len(e)
        row['flagged_evidence_rows']=int((e.qc_flags!='').sum())
        row['fishbase_version']=VERSION
        wide.append(row)
    wide=pd.DataFrame(wide)
    wide.to_csv(OUT/'life_history_80_species.csv',index=False)
    descriptions={
        'species':'Frozen project species key; existing subspecies pooling retained.',
        'investigation_tier':'Prior historical evidence screen; not a trait-based ranking.',
        'fishbase_name':'Source name used to attach data; blank when transfer withheld.',
        'fishbase_spec_code':'Attached FishBase identifier; blank when transfer withheld.',
        'taxon_match_status':'See taxon_crosswalk.csv for explicit source mapping and reasoning.',
        'positive_habitat_tags':'Only positive source flags; blanks/unchecked flags do not establish absence.',
        'positive_water_types':'Only positive Fresh/Brack/Saltwater source flags.',
        'maturity_length_n_records':'Count of maturity.Lm source rows including flagged rows; not independent studies.',
        'maturity_length_records':'Reported maturity lengths in cm per source record, retaining TL/SL/FL or unspecified type; sex and locality in Evidence. Do not average different length conventions.',
        'fecundity_n_records':'Count of rows with absolute fecundity values; not independent studies.',
        'fecundity_records':'Per-record absolute counts of eggs/offspring per female; time/batch basis preserved and unspecified basis flagged. Never sum records.',
        'larval_duration_n_records':'Count of source rows with duration values, including flagged rows; raw values in evidence.',
        'larval_duration_status':'Needs review when duration records exist; neither dispersal distance nor validated pelagic larval duration.',
        'generation_time_status':'Explicit estimate-versus-missing/taxonomy-withheld status; no observed generation times extracted.',
        'primary_source_review':'Database extraction only; cited original studies not independently verified.',
        'regional_applicability':'Source values have not been matched to proposal sites or historical eras.',
        'evidence_rows':'Number of nonempty extracted source-field values, including flags.',
        'flagged_evidence_rows':'Number of evidence rows requiring QC review.',
        'fishbase_version':'Pinned provider release label; not guaranteed database cutoff date.'}
    for field in wide:
        if field not in dictionary:
            describe(field,'metadata',descriptions[field])
    wide_dict=pd.DataFrame([dictionary[k] for k in wide])
    wide_dict.to_csv(OUT/'table_dictionary.csv',index=False)
    coverage=[]
    for s in FIELDS:
        e=ev[ev.trait==s['trait']]
        coverage.append(dict(trait=s['trait'],unit=s['unit'],evidence_type=s['evidence_type'],
                             species_with_value=e.species.nunique(),
                             species_with_unflagged_value=e.loc[e.qc_flags=='','species'].nunique(),
                             total_project_species=80,source_records=len(e)))
    coverage=pd.DataFrame(coverage)
    coverage.to_csv(OUT/'trait_coverage.csv',index=False)
    reviews=[]
    for r in cross.itertuples():
        if r.match_status!='exact_accepted_name':
            reviews.append(dict(species=r.species,reason=r.match_status,evidence_id='',detail=r.decision_basis))
    for r in ev[ev.qc_flags!=''].itertuples():
        reviews.append(dict(species=r.species,reason=r.qc_flags,evidence_id=r.evidence_id,detail='Source value preserved; review before use.'))
    for r in wide.itertuples():
        if not present(r.generation_time_model_years):
            reviews.append(dict(species=r.species,reason='generation_time_unavailable',evidence_id='',detail='Targeted literature review needed; do not substitute maturity age.'))
    have_refs=set(refs.RefNo.astype(int).astype(str))
    for r in ev.itertuples():
        if r.selected_reference_id and r.selected_reference_id not in have_refs:
            reviews.append(dict(species=r.species,reason='reference_id_not_in_snapshot',evidence_id=r.evidence_id,detail=r.selected_reference_id))
    reviews=pd.DataFrame(reviews)
    reviews.to_csv(OUT/'review_queue.csv',index=False)
    readme=pd.DataFrame([
        ('Purpose','Discovery table for 80 curated candidate species; not final analysis-ready traits.'),
        ('Source','FishBase team; version 26.06 via Source Cooperative / Carl Boettiger; CC-BY-NC-4.0.'),
        ('Generation time','FishBase Key Facts model estimates; not observed generation lengths. Bounds are not assumed 95% CI.'),
        ('Reported values','Reported by FishBase; underlying studies not independently reviewed. Population/sex ranges are not CIs.'),
        ('Missing','Blank means unavailable; counts of zero mean no extracted records, not biological zero.'),
        ('Taxonomy','Preserves project keys and existing subspecies pooling. 75 exact matches, 4 explicit aliases, 1 uncertain synonym withheld.'),
        ('Fecundity','Preserves record-specific basis; annual/batch/unspecified fecundities must not be pooled.'),
        ('Larval duration','Three source records have reversed min/max and unverified units; retained in Evidence and Review only.'),
        ('Use','Check Evidence, References and Review; establish regional relevance and verify primary sources before species selection.'),
        ('Rebuild','python3 Scripts/build_life_history.py; see docs/life_history_workflow.md.')],columns=['topic','note'])
    with pd.ExcelWriter(OUT/'life_history_80_species.xlsx',engine='openpyxl') as writer:
        sheets={'Read me':readme,'80 species':wide,'Coverage':coverage,'Taxon crosswalk':cross,
                'Evidence':ev,'References':refs,'Review':reviews,'Table dictionary':wide_dict,
                'Trait dictionary':pd.DataFrame(FIELDS)}
        for title,frame in sheets.items():
            # Excel/XML cannot represent some characters present in source
            # citations. Escape these visibly; authoritative CSV stays verbatim.
            def excel_value(value):
                if isinstance(value,str):
                    value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]',
                        lambda m:f'\\u{ord(m.group()):04x}',value)
                    if len(value)>32767:
                        raise ValueError('Excel cell would truncate; consult CSV and revise workbook layout')
                return value
            frame.map(excel_value).to_excel(writer,sheet_name=title,index=False)
            ws=writer.sheets[title]
            ws.freeze_panes='B2'
            ws.auto_filter.ref=ws.dimensions
            for cells in ws.iter_cols():
                letter=cells[0].column_letter
                ws.column_dimensions[letter].width=min(55,max(16,len(str(cells[0].value))+2))
            for cell in ws[1]:
                cell.font=openpyxl.styles.Font(bold=True,color='FFFFFF')
                cell.fill=openpyxl.styles.PatternFill('solid',fgColor='254B63')
            for cells in ws.iter_rows(min_row=2):
                for cell in cells:
                    if cell.data_type=='f': cell.data_type='s'
                    if isinstance(cell.value,float): cell.number_format='0.###'
    inputs=[ROSTER,CROSSWALK,RAW/'download_manifest.json',Path(__file__).resolve(),
            ROOT/'Scripts/download_fishbase_traits.py',ROOT/'requirements-life-history.txt',
            ROOT/'Data/All_collections.csv', ROOT/'docs/life_history_workflow.md',
            ROOT/'tests/test_life_history.py']
    # The independent validation report refers to this manifest; including its
    # hash here would introduce a circular dependency on the next validation run.
    outputs=[p for p in sorted(OUT.rglob('*')) if p.is_file()
             and p.name not in {'build_manifest.json','validation_report.json'}]
    revision=subprocess.run(['git','rev-parse','HEAD'],cwd=ROOT,capture_output=True,text=True,check=True).stdout.strip()
    status=subprocess.run(['git','status','--porcelain'],cwd=ROOT,capture_output=True,text=True,check=True).stdout.strip()
    manifest=dict(fishbase_version=VERSION,command='python3 Scripts/build_life_history.py',
                  git_head=revision,git_worktree_dirty=bool(status),
                  software=dict(python=platform.python_version(),numpy=np.__version__,pandas=pd.__version__,pyarrow=pyarrow.__version__,openpyxl=openpyxl.__version__),
                  inputs={str(p.relative_to(ROOT)):sha(p) for p in inputs},
                  raw_snapshots=download,outputs={str(p.relative_to(ROOT)):sha(p) for p in outputs},
                  counts=dict(species=len(wide),attached_taxa=len(codes),evidence_rows=len(ev),references=len(refs),review_rows=len(reviews)),
                  reproducibility='CSV content deterministic in pinned environment; XLSX zip timestamps and build metadata may differ.')
    (OUT/'build_manifest.json').write_text(json.dumps(manifest,indent=2,sort_keys=True)+'\n')
    print(json.dumps(manifest['counts'],indent=2))
    print(coverage[coverage.trait.isin(['generation_time_model','maximum_age_wild','maturity_age_reported','growth_k_reported','ReproMode','FecundityMin','LarvalDurationMin'])].to_string(index=False))


if __name__=='__main__':
    run()
